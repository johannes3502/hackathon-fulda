"""Source upload, extraction, chunking and retrieval. Synthetic data only."""

import io
import os
from pathlib import Path
from unittest.mock import AsyncMock, Mock
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.sources import db, ingest, retrieval
from app.sources import router as sources_router
from app.sources.chunking import chunk_segments
from app.sources.extract import ExtractionError, Segment, extract
from app.sources.models import Location, SourceHit
from app.sources.quotes import quote_match, verified_quote


def minimal_pdf(pages: list[str]) -> bytes:
    """Uncompressed PDF with one Helvetica text line per page."""
    count = len(pages)
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [" + b" ".join(f"{4 + 2 * i} 0 R".encode() for i in range(count))
               + f"] /Count {count} >>".encode(),
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >>"
                       f" /Contents {5 + 2 * i} 0 R >>".encode())
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def test_extract_pdf_pages(tmp_path):
    path = tmp_path / "quelle.pdf"
    path.write_bytes(minimal_pdf(["Umsatz 2025: 4,2 Mio. Euro", "", "Team: fuenf Personen"]))
    assert extract(path, ".pdf") == [Segment("page", 1, "Umsatz 2025: 4,2 Mio. Euro"),
                                     Segment("page", 3, "Team: fuenf Personen")]


def test_extract_pptx_slides_tables_and_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Marktgröße"
    first.placeholders[1].text = "30.000 Studierende in Hessen"
    first.notes_slide.notes_text_frame.text = "Quelle: synthetische Umfrage"
    deck.slides.add_slide(deck.slide_layouts[6])  # Empty slide is skipped.
    third = deck.slides.add_slide(deck.slide_layouts[6])
    table = third.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    for (row, col), text in {(0, 0): "Jahr", (0, 1): "Nutzer", (1, 0): "2026", (1, 1): "1200"}.items():
        table.cell(row, col).text = text
    path = tmp_path / "deck.pptx"
    deck.save(path)
    assert extract(path, ".pptx") == [
        Segment("slide", 1, "Marktgröße\n30.000 Studierende in Hessen\nNotizen: Quelle: synthetische Umfrage"),
        Segment("slide", 3, "Jahr | Nutzer\n2026 | 1200"),
    ]


def test_extract_docx_and_text(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Erster Absatz.")
    document.add_paragraph("")
    document.add_paragraph("Dritter   Absatz.")
    path = tmp_path / "notizen.docx"
    document.save(path)
    assert extract(path, ".docx") == [Segment("paragraph", 1, "Erster Absatz."), Segment("paragraph", 3, "Dritter Absatz.")]

    text = tmp_path / "notizen.md"
    text.write_bytes("# Titel\nZeile zwei\n\n\nAbsatz ab Zeile fünf\n".encode("cp1252"))
    assert extract(text, ".md") == [Segment("line", 1, "# Titel\nZeile zwei"), Segment("line", 5, "Absatz ab Zeile fünf")]


def test_extract_docx_keeps_tables_in_document_order(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Vor der Tabelle.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Kennzahl"
    table.cell(0, 1).text = "Wert"
    table.cell(1, 0).text = "Haushalte pro Woche"
    table.cell(1, 1).text = "1.200"
    document.add_paragraph("Nach der Tabelle.")
    path = tmp_path / "tabelle.docx"
    document.save(path)

    assert extract(path, ".docx") == [
        Segment("paragraph", 1, "Vor der Tabelle."),
        Segment("table", 1, "Kennzahl | Wert\nHaushalte pro Woche | 1.200"),
        Segment("paragraph", 2, "Nach der Tabelle."),
    ]


def test_chunking_does_not_merge_paragraphs_across_a_table():
    chunks = chunk_segments([
        Segment("paragraph", 1, "Kurz."),
        Segment("table", 1, "Kennzahl | Wert"),
        Segment("paragraph", 2, "Auch kurz."),
    ], max_chars=300)
    assert [(chunk.kind, chunk.start, chunk.end) for chunk in chunks] == [
        ("paragraph", 1, 1), ("table", 1, 1), ("paragraph", 2, 2)]


def test_short_exact_quotes_are_expanded_to_sentence_context():
    passage = "Der Preis beträgt 149 € statt 299 €. Die Aktion endet am Freitag."
    assert verified_quote("149 €", passage, 12) == ("Der Preis beträgt 149 € statt 299 €.", "exact")
    assert verified_quote("149", "Preis: 149 Euro.", 12) == ("Preis: 149 Euro.", "exact")
    assert verified_quote("149", "Preis: 1490 Euro.", 12) is None  # Not a substring of a longer number.
    assert verified_quote("5", "Wir haben 5 Standorte.", 12) is None
    assert verified_quote("%", "Wir haben 5 % Wachstum.", 12) is None
    assert verified_quote("2 Wochen", passage, 12) is None


@pytest.mark.parametrize("quote,passage", [
    ("Preis beträgt 149", "Preis beträgt 1490 Euro."),
    ("Preis beträgt 1.200", "Preis beträgt 1.200,50 Euro."),
    ("Preis beträgt 1.200", "Preis beträgt 1.200.000 Euro."),
    ("Der Umsatz beträgt 149", "Der Umsatz beträgt 149 Mio. Euro."),
    ("149 Euro", "Der Preis beträgt -149 Euro."),
    ("Wir haben 20 Nutzer", "Wir haben 20 Nutzern geholfen."),
    ("Die App hat viele aktive Nutzer und ist kostenlos.",
     "Die App hat viele aktive Nutzer und ist nicht, kostenlos."),
    ("Die App hat viele aktive Nutzer und ist kostenlos.",
     "Die App hat viele aktive Nutzer und ist niemals kostenlos."),
    ("Der Preis beträgt heute 149 Euro und bleibt dauerhaft günstig.",
     "Der Preis beträgt heute statt 149 Euro und bleibt dauerhaft günstig."),
])
def test_quote_verification_rejects_partial_words_and_changed_negation(quote, passage):
    assert verified_quote(quote, passage, 12) is None
    assert quote_match(quote, passage) is None


@pytest.mark.parametrize("quote,passage,minimum,expected", [
    ("149", "Preis: 149 Euro.", 1, "149"),
    ("AA BB", "Ein langer Satz endet AA\nBB beginnt den nächsten Satz.", 12,
     "Ein langer Satz endet AA\nBB beginnt den nächsten Satz."),
    ("Die Wirkungsweise ist großartig.", "Die Wirkungs-\nweise ist großartig.", 12,
     "Die Wirkungs-\nweise ist großartig."),
    ("Die Straße ist lang genug.", "Die STRASSE ist lang genug.", 12,
     "Die STRASSE ist lang genug."),
    ('Die App ist "kostenlos" verfügbar.', 'Die App ist „kostenlos“ verfügbar.', 12,
     'Die App ist „kostenlos“ verfügbar.'),
    ("Diese lange Quelle enthält mehrere Wörter und einen Schreibfeler im wichtigen mittleren Satz.",
     "Diese lange Quelle enthält mehrere Wörter und einen Schreibfehler im wichtigen mittleren Satz.", 12,
     "Diese lange Quelle enthält mehrere Wörter und einen Schreibfehler im wichtigen mittleren Satz."),
])
def test_verified_quotes_always_return_the_complete_source_span(quote, passage, minimum, expected):
    found = verified_quote(quote, passage, minimum)
    assert found is not None and found[0] == expected and found[0] in passage


def test_multi_query_fusion_prefers_passages_found_by_several_queries():
    def hit(chunk_id: int, text: str, similarity: float) -> SourceHit:
        return SourceHit(chunk_id=chunk_id, source_id=uuid4(), filename="folien.pdf",
                         location=Location(kind="page", start=1, end=1), text=text,
                         similarity=similarity, matched_by=["vector"])

    first, shared, last = hit(1, "A", 0.8), hit(2, "B", 0.4), hit(3, "C", 0.7)
    fused = retrieval.fuse([[first, shared], [shared, last]], k=2)
    assert [item.chunk_id for item in fused] == [2, 1]
    assert fused[0].retrieval_score > fused[1].retrieval_score


def test_extract_rejects_broken_files(tmp_path):
    path = tmp_path / "kaputt.pdf"
    path.write_bytes(b"not a pdf")
    with pytest.raises(ExtractionError):
        extract(path, ".pdf")


def test_chunking_keeps_pages_apart_and_merges_small_paragraphs():
    sentence = "Dies ist ein synthetischer Satz mit einigen Wörtern."  # 52 characters
    long_page = Segment("page", 2, " ".join([sentence] * 40))
    chunks = chunk_segments([Segment("page", 1, "Kurz."), long_page], max_chars=300, overlap_chars=60)
    assert chunks[0].text == "Kurz." and (chunks[0].start, chunks[0].end) == (1, 1)
    rest = chunks[1:]
    assert len(rest) > 1 and all(c.kind == "page" and c.start == c.end == 2 for c in rest)
    assert all(len(c.text) <= 300 for c in rest)
    assert rest[1].text.startswith(sentence)  # Overlap carries the previous sentence.
    assert [c.ordinal for c in chunks] == list(range(len(chunks)))

    merged = chunk_segments([Segment("paragraph", n, f"Absatz {n}.") for n in (1, 2, 4)], max_chars=300)
    assert len(merged) == 1 and (merged[0].start, merged[0].end) == (1, 4)


def test_chunking_splits_overlong_sentences():
    chunks = chunk_segments([Segment("line", 1, "wort " * 200)], max_chars=100, overlap_chars=0)
    assert len(chunks) > 5 and all(len(c.text) <= 100 for c in chunks)


def test_quote_verification_and_labels():
    chunk = "Laut  Umfrage nutzen „30.000 Studierende“ die App –\nwöchentlich."
    assert retrieval.quote_in_text('"30.000 studierende" die App - wöchentlich', chunk)
    assert not retrieval.quote_in_text("40.000 Studierende", chunk)
    assert not retrieval.quote_in_text(" … ", chunk)
    assert retrieval.location_label(Location(kind="slide", start=2, end=2)) == "Folie 2"
    assert retrieval.location_label(Location(kind="paragraph", start=1, end=4)) == "Absätze 1–4"
    assert retrieval.location_label(Location(kind="table", start=1, end=1)) == "Tabelle 1"


def test_upload_validation_and_accept(monkeypatch):
    client = TestClient(app)
    collection = str(uuid4())
    response = client.post("/sources", data={"collection_id": collection}, files={"file": ("x.exe", b"MZ")})
    assert response.status_code == 415
    response = client.post("/sources", data={"collection_id": collection}, files={"file": ("x.txt", b"")})
    assert response.status_code == 422
    monkeypatch.setenv("SOURCE_MAX_BYTES", "4")
    response = client.post("/sources", data={"collection_id": collection}, files={"file": ("x.txt", b"12345")})
    assert response.status_code == 413
    monkeypatch.delenv("SOURCE_MAX_BYTES")

    store, enqueue = Mock(), AsyncMock(return_value="task")
    monkeypatch.setattr(sources_router, "_store", store)
    monkeypatch.setattr(sources_router, "enqueue_task", enqueue)
    response = client.post("/sources", data={"collection_id": collection},
                           files={"file": ("C:\\Users\\demo\\Folien.PPTX", b"synthetic", "application/octet-stream")})
    assert response.status_code == 202
    body = response.json()
    assert body["source"]["filename"] == "Folien.PPTX" and body["source"]["status"] == "queued"
    assert body["status_url"] == f"/sources/{body['source']['id']}"
    enqueue.assert_awaited_once()
    assert enqueue.await_args.args[1] == body["source"]["id"]


def test_ingest_marks_empty_documents_failed(monkeypatch, tmp_path):
    from app.sources.models import Source

    source = Source(id=uuid4(), collection_id=uuid4(), filename="leer.txt", content_type="text/plain",
                    size_bytes=3, status="queued", created_at="2026-10-24T12:00:00Z")
    statuses = []
    monkeypatch.setattr(db, "get_source", lambda source_id: (source, "key"))
    monkeypatch.setattr(db, "set_status", lambda source_id, status, error=None: statuses.append((status, error)))
    s3 = Mock()
    s3.download_file.side_effect = lambda bucket, key, path: Path(path).write_text("\n \n")
    monkeypatch.setattr(ingest, "s3_client", lambda: s3)
    monkeypatch.setenv("S3_BUCKET", "test")
    ingest.ingest_source(source.id)
    assert statuses[0] == ("processing", None)
    assert statuses[1][0] == "failed" and "Keine Textinhalte" in statuses[1][1]


class FakeEmbedder:
    """Deterministic bag-of-words vectors; only for the database contract test."""

    model = "test-bag-of-words"
    vocabulary = ["umsatz", "nutzer", "app", "team", "studierende", "euro"]

    def _vector(self, text):
        words = text.lower()
        vector = [float(words.count(term)) for term in self.vocabulary] + [0.01]
        norm = sum(v * v for v in vector) ** 0.5
        return [v / norm for v in vector]

    def embed_documents(self, texts):
        return [self._vector(text) for text in texts]

    def embed_query(self, text):
        return self._vector(text)


@pytest.mark.skipif(not os.getenv("POSTGRES_HOST"), reason="Needs PostgreSQL with pgvector (docker compose)")
def test_database_hybrid_search_roundtrip():
    from app.sources.models import Source

    collection, other = uuid4(), uuid4()
    embedder = FakeEmbedder()
    ids = []
    try:
        for collection_id, texts in ((collection, ["Der Umsatz stieg auf 4,2 Mio. Euro.", "Die App hat 1200 Nutzer."]),
                                     (other, ["Die App hat 99 Nutzer."])):
            source = Source(id=uuid4(), collection_id=collection_id, filename="synthetisch.txt", content_type="text/plain",
                            size_bytes=1, status="queued", created_at="2026-10-24T12:00:00Z")
            ids.append(source.id)
            db.insert_source(source, f"sources/{source.id}/original")
            chunks = chunk_segments([Segment("line", n, text) for n, text in enumerate(texts, 1)], max_chars=40)
            assert db.store_chunks(source.id, chunks, embedder.embed_documents([c.text for c in chunks]), embedder.model)
            assert db.store_chunks(source.id, chunks, embedder.embed_documents([c.text for c in chunks]), embedder.model)

        hits = retrieval.search(collection, "Wie viele Nutzer hat die App?", k=2, embedder=embedder)
        assert [hit.text for hit in hits][0] == "Die App hat 1200 Nutzer."
        assert hits[0].matched_by == ["vector", "keyword"] and hits[0].similarity > 0.5
        assert all(hit.source_id == ids[0] for hit in hits)
        # Keyword-only mode still finds exact figures.
        keyword_hits = retrieval.search(collection, "Wie hoch war der Umsatz?", k=5, embedder=None)
        assert [(hit.text, hit.matched_by, hit.similarity) for hit in keyword_hits] == [
            ("Der Umsatz stieg auf 4,2 Mio. Euro.", ["keyword"], None)]
        assert [hit.text for hit in retrieval.search(collection, "4,2", embedder=None)] == [
            "Der Umsatz stieg auf 4,2 Mio. Euro."]
        assert retrieval.search(collection, "und der die", embedder=None) == []
        assert db.list_sources(collection)[0].chunk_count == 2
    finally:
        for source_id in ids:
            db.delete_source(source_id)
    assert db.get_source(ids[0]) is None
