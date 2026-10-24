"""Benchmark the existing source pipeline with synthetic German slides.

Run from backend: python -u -m scripts.benchmark_vectors
Set EMBEDDING_MODEL_PATH to the pinned, downloaded model. Downloads are not
timed. Set the usual POSTGRES_* variables to include storage and hybrid search.
Only this run's UUID sources are created/deleted; existing sources are untouched.
"""

import json
import math
import os
import platform
import resource
import statistics
import tempfile
import time
from datetime import UTC, datetime
from dataclasses import replace
from importlib.metadata import version
from pathlib import Path
from uuid import uuid4

from pptx import Presentation

from app.sources.chunking import chunk_segments
from app.sources.embeddings import configured_embedder
from app.sources.extract import extract

QUERIES = [
    "Wie unterstützt das Projekt Studierende beim Präsentieren?",
    "Wie werden Quellen und Zitate geprüft?",
    "Welche Daten bleiben nach der Aufnahme gespeichert?",
    "Wie erhält das Team Rückmeldungen zum Prototyp?",
    "Wie viele Personen nehmen an der synthetischen Übung teil?",
]
BODY = (
    "Das fiktive Lernprojekt unterstützt Studierende beim Üben von Präsentationen. "
    "Eine synthetische Übungsgruppe besteht aus zwölf Personen und trifft sich wöchentlich. "
    "Die Teilnehmenden laden ihre Quellen vor dem Vortrag hoch und besprechen anschließend offene Fragen. "
    "Das System zerlegt Dokumente in Textabschnitte und bewahrt die zugehörige Foliennummer auf. "
    "Eine Suche verbindet ähnliche Formulierungen mit exakten Begriffen und Zahlen aus den Quellen. "
    "Rückmeldungen verweisen auf konkrete Stellen, damit Aussagen und Zitate nachvollziehbar bleiben. "
    "Die Originalaufnahme wird nach der Verarbeitung gelöscht; kurze Ausschnitte dienen der Besprechung. "
    "Das Team überprüft den Prototyp anhand vorbereiteter Beispiele und behauptet keine validierte Genauigkeit. "
    "Alle Namen, Zahlen und Situationen in diesem Dokument sind erfunden und dienen ausschließlich dem Leistungstest."
)


def timed(fn):
    started = time.perf_counter()
    value = fn()
    return value, (time.perf_counter() - started) * 1000


def emit(stage, **values):
    print(json.dumps({"stage": stage, **values}, ensure_ascii=False), flush=True)


def summary(samples):
    ordered = sorted(samples)
    return {"n": len(samples), "median_ms": statistics.median(samples),
            "p95_ms": ordered[math.ceil(len(samples) * .95) - 1],
            "min_ms": min(samples), "max_ms": max(samples), "samples_ms": samples}


def validate(vectors, count):
    assert len(vectors) == count
    for vector in vectors:
        assert len(vector) == 768 and all(math.isfinite(x) for x in vector)
        assert abs(sum(x * x for x in vector) - 1) < .01


def main():
    assert os.getenv("EMBEDDING_MODEL_PATH"), "Download the pinned model first; set EMBEDDING_MODEL_PATH"
    emit("environment", python=platform.python_version(), platform=platform.platform(),
         packages={name: version(name) for name in ("fastembed", "onnxruntime", "numpy", "python-pptx")})
    embedder, elapsed = timed(configured_embedder)
    assert embedder is not None, "Enable local embeddings"
    emit("model_load", ms=elapsed, model=embedder.model,
         providers=embedder._model.model.model.get_providers())
    vector, elapsed = timed(lambda: embedder.embed_query(QUERIES[0]))
    validate([vector], 1)
    emit("first_query", ms=elapsed)
    last_chunks, last_vectors = None, None
    with tempfile.TemporaryDirectory(prefix="vector-benchmark-") as workdir:
        for count in (3, 20, 100):
            deck = Presentation()
            for index in range(count):
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                slide.shapes.title.text = f"Synthetische Übung {index + 1}"
                slide.placeholders[1].text = BODY
            path = Path(workdir) / f"synthetic-{count}.pptx"
            deck.save(path)
            for repeat in range(3):
                segments, extraction_ms = timed(lambda: extract(path, ".pptx"))
                chunks, chunking_ms = timed(lambda: chunk_segments(segments))
                assert len(segments) == len(chunks) == count
                vectors, embedding_ms = timed(lambda: embedder.embed_documents([c.text for c in chunks]))
                validate(vectors, count)
                emit("ingest_compute", slides=count, chunks=len(chunks), repeat=repeat + 1,
                     mean_chars=statistics.mean(len(c.text) for c in chunks),
                     extraction_ms=extraction_ms, chunking_ms=chunking_ms, embedding_ms=embedding_ms,
                     total_ms=extraction_ms + chunking_ms + embedding_ms,
                     chunks_per_second=count / (embedding_ms / 1000))
            last_chunks, last_vectors = chunks, vectors
    samples = []
    for query in QUERIES * 4:
        vector, elapsed = timed(lambda: embedder.embed_query(query))
        validate([vector], 1)
        samples.append(elapsed)
    emit("warm_query_embedding", **summary(samples))

    if os.getenv("POSTGRES_HOST"):
        from app.sources import db, retrieval
        from app.sources.models import Source

        source = Source(id=uuid4(), collection_id=uuid4(), filename="synthetic-benchmark.pptx",
                        content_type="application/octet-stream", size_bytes=1, status="queued",
                        created_at=datetime.now(UTC))
        try:
            db.insert_source(source, f"benchmark/{source.id}")
            with db.connect() as connection:
                emit("database", postgres=connection.execute("SELECT version()").fetchone()["version"],
                     pgvector=connection.execute("SELECT extversion FROM pg_extension WHERE extname = 'vector'")
                     .fetchone()["extversion"])
            for count in (100, 1000, 10000):
                # Replication measures DB scaling, not embedding throughput or relevance.
                chunks = [replace(last_chunks[i % 100], ordinal=i) for i in range(count)]
                vectors = [last_vectors[i % 100] for i in range(count)]
                stored, elapsed = timed(lambda: db.store_chunks(source.id, chunks, vectors, embedder.model))
                assert stored and db.get_source(source.id)[0].chunk_count == count
                emit("store_chunks", chunks=count, ms=elapsed)
                for mode, model in (("hybrid", embedder), ("keyword_only", None)):
                    samples = []
                    for query in QUERIES * 4:
                        hits, elapsed = timed(lambda: retrieval.search(source.collection_id, query, embedder=model))
                        assert hits and all(hit.source_id == source.id for hit in hits)
                        if model:
                            assert any("vector" in hit.matched_by for hit in hits)
                        samples.append(elapsed)
                    emit(mode, chunks=count, **summary(samples))
        finally:
            db.delete_source(source.id)
        assert db.get_source(source.id) is None
    else:
        emit("database_skipped", reason="POSTGRES_HOST is not configured")
    emit("peak_process_rss", bytes=resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
         * (1 if platform.system() == "Darwin" else 1024))


if __name__ == "__main__":
    main()
