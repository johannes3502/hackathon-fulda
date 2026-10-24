"""Exercise source upload → Redis → worker → embeddings → hybrid search → delete.

Uses a synthetic PPTX built in memory; run inside the backend container.
"""

import io
import time
from uuid import uuid4

import httpx
from pptx import Presentation

SLIDES = [
    ("Problem", "Studierende üben Präsentationen selten mit Feedback."),
    ("Markt", "In Hessen studieren laut synthetischer Annahme 260.000 Menschen."),
    ("Team", "Fünf Personen aus Informatik und Design bauen den Prototyp."),
]


def deck() -> bytes:
    presentation = Presentation()
    for title, body in SLIDES:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def main():
    collection = str(uuid4())
    source_id = None
    with httpx.Client(base_url="http://localhost:8000", timeout=60) as client:
        try:
            response = client.post("/sources", data={"collection_id": collection},
                                   files={"file": ("synthetisch.pptx", deck())})
            assert response.status_code == 202, response.text
            accepted = response.json()
            source_id = accepted["source"]["id"]
            deadline = time.monotonic() + 120
            while time.monotonic() < deadline:
                source = client.get(accepted["status_url"]).raise_for_status().json()
                if source["status"] in ("ready", "failed"):
                    break
                time.sleep(0.5)
            assert source["status"] == "ready", source.get("error") or source["status"]
            assert source["chunk_count"] == len(SLIDES) and source["embedding_model"]

            started = time.monotonic()
            hits = client.post("/sources/search", json={
                "collection_id": collection, "query": "Wer arbeitet am Projekt mit?", "k": 2,
            }).raise_for_status().json()
            assert hits[0]["location"] == {"kind": "slide", "start": 3, "end": 3}, hits
            assert "vector" in hits[0]["matched_by"]
            exact = client.post("/sources/search", json={
                "collection_id": collection, "query": "260.000", "k": 1,
            }).raise_for_status().json()
            assert exact[0]["location"]["start"] == 2 and "keyword" in exact[0]["matched_by"], exact
            listed = client.get("/sources", params={"collection_id": collection}).raise_for_status().json()
            assert [item["id"] for item in listed] == [source_id]
            print(f"Sources smoke passed: {source['embedding_model']}, "
                  f"search {time.monotonic() - started:.2f}s, top similarity {hits[0]['similarity']}")
        finally:
            if source_id:
                assert client.delete(f"/sources/{source_id}").status_code in (204, 404)
        assert client.get(f"/sources/{source_id}").status_code == 404


if __name__ == "__main__":
    main()
