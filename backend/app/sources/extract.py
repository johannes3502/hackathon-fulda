"""Text per page (PDF), slide (PPTX), paragraph/table (DOCX) or line block (TXT/MD).

Every segment keeps its 1-based location so answers can cite "Seite 3" or
"Folie 2". No OCR: scanned PDFs without a text layer yield no segments.
"""

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

LocationKind = Literal["page", "slide", "paragraph", "table", "line"]

SUPPORTED: dict[str, LocationKind] = {
    ".pdf": "page", ".pptx": "slide", ".docx": "paragraph", ".txt": "line", ".md": "line",
}


@dataclass(frozen=True)
class Segment:
    kind: LocationKind
    number: int
    text: str


class ExtractionError(ValueError):
    """The file is not readable; the message is safe to show to the user."""


def clean(text: str) -> str:
    """Collapse spaces per line and drop empty lines; keep line structure."""
    text = text.replace("­", "").replace("\x00", "")
    lines = (re.sub(r"[ \t ​]+", " ", line).strip() for line in text.splitlines())
    return "\n".join(line for line in lines if line)


def extract(path: Path, suffix: str) -> list[Segment]:
    readers = {".pdf": _pdf, ".pptx": _pptx, ".docx": _docx, ".txt": _text, ".md": _text}
    if suffix not in readers:
        raise ExtractionError("Dateityp wird nicht unterstützt.")
    try:
        segments = readers[suffix](path)
    except ExtractionError:
        raise
    except Exception as error:  # Parser libraries raise many unrelated types.
        raise ExtractionError("Datei konnte nicht gelesen werden. Ist sie beschädigt?") from error
    return [segment for segment in segments if segment.text]


def _pdf(path: Path) -> list[Segment]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ExtractionError("Passwortgeschützte PDFs werden nicht unterstützt.")
    return [Segment("page", number, clean(page.extract_text() or "")) for number, page in enumerate(reader.pages, 1)]


def _shape_texts(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from _shape_texts(shape.shapes)
        elif shape.has_text_frame:
            yield shape.text_frame.text
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                yield " | ".join(cell.text for cell in row.cells)


def _pptx(path: Path) -> list[Segment]:
    from pptx import Presentation

    segments = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        parts = list(_shape_texts(slide.shapes))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notizen: {notes}")
        segments.append(Segment("slide", number, clean("\n".join(parts))))
    return segments


def _docx(path: Path) -> list[Segment]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    segments: list[Segment] = []
    paragraph_number = table_number = 0
    # `iter_inner_content` keeps paragraphs and tables in document order.
    # Paragraphs inside a table are represented by the table's rows.
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_number += 1
            segments.append(Segment("paragraph", paragraph_number, clean(block.text)))
        elif isinstance(block, Table):
            table_number += 1
            rows = []
            for row in block.rows:
                cells = [clean(cell.text).replace("\n", " ").strip() for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            segments.append(Segment("table", table_number, clean("\n".join(rows))))
    return segments


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Textdatei ist weder UTF-8 noch Windows-1252.")


def _text(path: Path) -> list[Segment]:
    """Blocks separated by blank lines; the number is the block's first line."""
    segments, block, start = [], [], 1
    for number, line in enumerate(_decode(path.read_bytes()).splitlines(), 1):
        if line.strip():
            if not block:
                start = number
            block.append(line)
        elif block:
            segments.append(Segment("line", start, clean("\n".join(block))))
            block = []
    if block:
        segments.append(Segment("line", start, clean("\n".join(block))))
    return segments
