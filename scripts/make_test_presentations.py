"""Generate synthetic test presentations for the source/RAG and review pipeline.

Creates three demos in ``test-presentations/``:
  01-startup-plantsens.pptx   10 slides, speaker notes, tables (startup pitch)
  02-thesis-kurzpausen.pdf    6 pages with a text layer (thesis defense)
  03-nonprofit-tafel-musterstadt.docx  paragraphs, table, FAQ (nonprofit)

Every document is fictional and uses synthetic demo data. Numbers are chosen so
the fact check can show all three verdicts (supports / contradicts / not found)
when the spoken script in ``sprech-skript.txt`` is used.

Run with the document toolchain (python-pptx, python-docx, reportlab):
    python scripts/make_test_presentations.py
"""

from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "test-presentations"

INK = "16241D"
PRIMARY = "146B45"
ACCENT = "4ADE80"
MUTED = "5F7268"
TINT = "EAF5EF"
DARK = "0B2B1F"
LIGHT = "D7F5E4"
WHITE = "FFFFFF"
FONT = "Calibri"


# --------------------------------------------------------------------------- PPTX

def _bg(slide, prs, color: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    shape.shadow.inherit = False


def _text(slide, x, y, w, h, text, size=18, bold=False, color=INK, align=None, anchor=None, spacing=1.0):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor if anchor is not None else MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align if align is not None else PP_ALIGN.LEFT
    paragraph.line_spacing = spacing
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = FONT
        run.font.color.rgb = RGBColor.from_string(color)
    return box


def _bullet(paragraph, color: str) -> None:
    from lxml import etree
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(Inches(0.28)))
    p_pr.set("indent", str(-Inches(0.28)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont"):
        for element in p_pr.findall(qn(tag)):
            p_pr.remove(element)
    bu_clr = etree.SubElement(p_pr, qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(p_pr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(p_pr, qn("a:buChar")).set("char", "•")


def _bullets(slide, x, y, w, h, items, size=19, gap=14, color=INK):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = item
        run.font.size = Pt(size)
        run.font.name = FONT
        run.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(gap)
        paragraph.line_spacing = 1.12
        _bullet(paragraph, ACCENT)
    return box


def _card(slide, x, y, w, h, value, label) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(TINT)
    shape.line.color.rgb = RGBColor.from_string(ACCENT)
    shape.line.width = Inches(0.015)
    shape.shadow.inherit = False
    _text(slide, x + 0.3, y + 0.35, w - 0.6, 1.2, value, size=44, bold=True, color=PRIMARY,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    _text(slide, x + 0.3, y + 1.65, w - 0.6, h - 1.9, label, size=15, color=MUTED, spacing=1.15)


def _table(slide, x, y, w, rows, widths, row_height=0.55, size=14):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
    from lxml import etree

    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(row_height * len(rows)))
    table = shape.table
    tbl = shape._element.graphic.graphicData.tbl
    tbl_pr = tbl.tblPr
    for element in tbl_pr.findall(qn("a:tableStyleId")):
        tbl_pr.remove(element)
    etree.SubElement(tbl_pr, qn("a:tableStyleId")).text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for row_index, row in enumerate(rows):
        table.rows[row_index].height = Inches(row_height)
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.margin_left = Inches(0.14)
            cell.margin_right = Inches(0.14)
            cell.margin_top = Inches(0.06)
            cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                PRIMARY if row_index == 0 else (TINT if row_index % 2 else WHITE)
            )
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(size)
                    run.font.name = FONT
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor.from_string(WHITE if row_index == 0 else INK)
    return table


def _slide(prs, eyebrow: str, title: str):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(0.55), Inches(0.55), Inches(0.055))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
    bar.line.fill.background()
    bar.shadow.inherit = False
    _text(slide, 0.85, 0.7, 8.0, 0.35, eyebrow.upper(), size=11, bold=True, color=PRIMARY)
    _text(slide, 0.85, 1.05, 11.4, 0.8, title, size=31, bold=True, color=INK)
    return slide


def build_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(cover, prs, DARK)
    bar = cover.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.05), Inches(0.8), Inches(0.07))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
    bar.line.fill.background()
    bar.shadow.inherit = False
    _text(cover, 0.9, 1.35, 10.0, 0.4, "PITCH DECK · SYNTHETISCHE DEMO", size=12, bold=True, color=ACCENT)
    _text(cover, 0.9, 1.95, 10.8, 2.2, "PlantSens", size=64, bold=True, color=WHITE)
    _text(cover, 0.9, 3.35, 10.6, 1.2, "Bewässerung nach Bedarf für kleine Gewächshäuser",
          size=26, color=LIGHT, spacing=1.1)
    _text(cover, 0.9, 5.15, 6.0, 0.4, "PlantSens GmbH · Kassel · Pilot 2025", size=14, color=LIGHT)
    _text(cover, 0.9, 5.55, 6.0, 0.4, "Stand: Oktober 2025", size=14, color="8FB9A0")
    _text(cover, 0.9, 6.65, 9.0, 0.4, "Synthetische Daten. Kein echtes Unternehmen, keine echten Personen.",
          size=11, color="8FB9A0")

    slide = _slide(prs, "Problem", "Kleine Betriebe gießen nach Zeitplan, nicht nach Bedarf")
    _bullets(slide, 0.85, 2.15, 7.9, 3.9, [
        "Kleine Gewächshäuser gießen nach festen Zeitplänen statt nach der Bodenfeuchte.",
        "In der Pilotbefragung gaben 12 von 14 Betrieben an, nach Gefühl zu gießen.",
        "Im Schnitt gehen 38 % des Wassers verloren.",
    ])
    _card(slide, 9.3, 1.95, 3.2, 3.4, "38 %", "Wasserverlust in kleinen Gewächshäusern ohne Messung")
    slide.notes_slide.notes_text_frame.text = (
        "Synthetische Pilotbefragung mit 14 Betrieben in Nordhessen, März 2025. "
        "Die 38 % sind eine Schätzung aus Betriebsangaben, keine Messreihe."
    )

    slide = _slide(prs, "Lösung", "Drei Sensoren, eine App, 27 % weniger Wasser")
    _bullets(slide, 0.85, 2.15, 7.9, 3.9, [
        "Drei Bodenfeuchtesensoren pro 100 m² messen alle zehn Minuten.",
        "Die App empfiehlt Bewässerung erst unter 35 % Bodenfeuchte.",
        "Im sechswöchigen Pilot sank der Wasserverbrauch um 27 %.",
    ])
    _card(slide, 9.3, 1.95, 3.2, 3.4, "27 %", "weniger Wasserverbrauch im sechswöchigen Pilot")
    slide.notes_slide.notes_text_frame.text = (
        "Die 27 % beziehen sich auf neun Pilotbetriebe mit je drei Sensoren über sechs Wochen. "
        "Die Zahl ist nicht unabhängig validiert."
    )

    slide = _slide(prs, "Markt", "4.200 Betriebe in Hessen, 15 % Zielanteil")
    _bullets(slide, 0.85, 2.15, 7.9, 3.9, [
        "4.200 Kleinbetriebe mit 100 bis 1.000 m² in Hessen.",
        "Zielmarkt sind Betriebe ohne feste Bewässerungsanlage.",
        "Geplanter Marktanteil: 15 % in drei Jahren.",
    ])
    _card(slide, 9.3, 1.95, 3.2, 3.4, "4.200", "Kleinbetriebe im Zielmarkt Hessen")
    slide.notes_slide.notes_text_frame.text = (
        "Eigene Auswertung von 14 Landkreisen. Keine externe Marktstudie, keine Kaufkraftanalyse."
    )

    slide = _slide(prs, "Geschäftsmodell", "Hardware plus Abo, amortisiert in 14 Monaten")
    _table(slide, 0.85, 2.0, 11.6, [
        ["Position", "Wert"],
        ["Hardware", "149 € einmalig"],
        ["Abo", "79 € pro Monat"],
        ["Sensoren pro Betrieb", "3 pro 100 m²"],
        ["Amortisation für den Kunden", "14 Monate"],
    ], [4.9, 6.7], row_height=0.72, size=16)
    slide.notes_slide.notes_text_frame.text = (
        "Die Amortisation ist aus Wasserkosten und Ertragsstabilisierung modelliert, nicht gemessen."
    )

    slide = _slide(prs, "Traction", "Neun Pilotbetriebe, alle zahlen")
    _bullets(slide, 0.85, 2.15, 7.9, 3.9, [
        "9 Pilotbetriebe, alle 9 zahlen das Abo.",
        "Monatlich wiederkehrender Umsatz: 711 €.",
        "Churn seit dem Start: 0.",
    ])
    _card(slide, 9.3, 1.95, 3.2, 3.4, "711 €", "MRR aus neun Abos zu je 79 €")
    slide.notes_slide.notes_text_frame.text = (
        "Stand Oktober 2025. 9 Abos zu 79 € ergeben 711 € MRR. Keine Investorenfinanzierung."
    )

    slide = _slide(prs, "Wettbewerb", "Günstiger und lokal beraten")
    _bullets(slide, 0.85, 2.15, 7.9, 3.9, [
        "Bestehende Systeme kosten ab 900 € Anschaffung.",
        "Kein Anbieter bietet lokale Beratung in Hessen.",
        "Unsere Installation dauert unter zwei Stunden.",
    ])
    _card(slide, 9.3, 1.95, 3.2, 3.4, "ab 900 €", "Anschaffung bei bestehenden Systemen")
    slide.notes_slide.notes_text_frame.text = (
        "Wettbewerbspreise aus öffentlichen Preislisten, Stand September 2025. "
        "Lokale Beratung ist eine Einschätzung aus zwölf Sondierungsgesprächen."
    )

    slide = _slide(prs, "Team", "Drei Personen, drei Disziplinen")
    _bullets(slide, 0.85, 2.15, 7.9, 3.9, [
        "Drei Personen im Kernteam, alle in Vollzeit.",
        "Agrarwissenschaft, IoT-Entwicklung und Vertrieb.",
        "Zwei Gründerinnen und ein Gründer.",
    ])
    _card(slide, 9.3, 1.95, 3.2, 3.4, "3", "Personen im Kernteam")
    slide.notes_slide.notes_text_frame.text = (
        "Keine Namen und keine Lebensläufe in dieser synthetischen Demo."
    )

    slide = _slide(prs, "Finanzierung", "250.000 € für 18 Monate")
    _bullets(slide, 0.85, 2.15, 7.9, 3.9, [
        "250.000 € für 18 Monate Laufzeit.",
        "60 % Produktentwicklung, 40 % Vertrieb.",
        "Meilenstein: 60 zahlende Betriebe bis Q4 2026.",
    ])
    _card(slide, 9.3, 1.95, 3.2, 3.4, "250.000 €", "Finanzierungsbedarf für 18 Monate")
    slide.notes_slide.notes_text_frame.text = (
        "Die Aufteilung ist ein Planwert. Es gibt noch keine Investorenzusagen."
    )

    slide = _slide(prs, "Roadmap", "Vom Pilot zur Serie in vier Quartalen")
    _table(slide, 0.85, 2.0, 11.6, [
        ["Quartal", "Meilenstein"],
        ["Q1 2026", "Sensor v2 mit Solarbetrieb"],
        ["Q2 2026", "Serienfertigung bei einem Partner in Hessen"],
        ["Q3 2026", "40 zahlende Betriebe"],
        ["Q4 2026", "60 zahlende Betriebe"],
    ], [2.6, 9.0], row_height=0.72, size=16)
    slide.notes_slide.notes_text_frame.text = "Alle Ziele sind Planwerte, keine Zusagen."

    prs.core_properties.title = "PlantSens — Bewässerung nach Bedarf"
    prs.core_properties.author = "Synthetische Demo"
    prs.core_properties.subject = "Testpräsentation für den Präsentationstrainer"
    prs.save(path)


# ---------------------------------------------------------------------------- PDF

def build_pdf(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (HRFlowable, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    green = colors.HexColor("#146B45")
    dark = colors.HexColor("#16241D")
    muted = colors.HexColor("#5F7268")
    accent = colors.HexColor("#4ADE80")
    tint = colors.HexColor("#EAF5EF")

    styles = {
        "label": ParagraphStyle("label", fontName="Helvetica-Bold", fontSize=9, leading=12,
                                textColor=green, spaceAfter=10),
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=24, leading=30,
                                textColor=dark, spaceAfter=14),
        "subtitle": ParagraphStyle("subtitle", fontName="Helvetica", fontSize=13, leading=18,
                                   textColor=muted, spaceAfter=6),
        "h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=15, leading=19,
                             textColor=green, spaceBefore=8, spaceAfter=10),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=10.5, leading=15,
                               textColor=dark, spaceAfter=9, alignment=TA_LEFT),
        "small": ParagraphStyle("small", fontName="Helvetica-Oblique", fontSize=9, leading=12,
                                textColor=muted, spaceAfter=4),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=8.5, leading=11, textColor=dark),
        "cell_bold": ParagraphStyle("cell_bold", fontName="Helvetica-Bold", fontSize=8.5, leading=11,
                                    textColor=colors.white),
    }

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setStrokeColor(accent)
        canvas.setLineWidth(2)
        canvas.line(2 * cm, 1.65 * cm, 3.2 * cm, 1.65 * cm)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(muted)
        canvas.drawString(2 * cm, 1.25 * cm, "Synthetische Beispielstudie — keine echten Personen")
        canvas.drawRightString(A4[0] - 2 * cm, 1.25 * cm, f"Seite {doc.page}")
        canvas.restoreState()

    def table(rows, widths, align_right=False) -> Table:
        data = [[Paragraph(str(cell), styles["cell_bold"] if index == 0 else styles["cell"])
                 for cell in row] for index, row in enumerate(rows)]
        result = Table(data, colWidths=widths, hAlign="LEFT")
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), green),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#C9DCD3")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ]
        for row_index in range(1, len(rows)):
            if row_index % 2 == 0:
                style.append(("BACKGROUND", (0, row_index), (-1, row_index), tint))
        if align_right:
            style.append(("ALIGN", (1, 0), (-1, -1), "RIGHT"))
        result.setStyle(TableStyle(style))
        return result

    doc = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2.3 * cm,
        title="Kurze Pausen, bessere Konzentration?", author="Synthetische Demo",
        subject="Testpräsentation für den Präsentationstrainer",
    )

    story = [
        Spacer(1, 3.2 * cm),
        Paragraph("SYNTHETISCHE BEISPIELSTUDIE", styles["label"]),
        Paragraph("Kurze Pausen, bessere Konzentration?", styles["title"]),
        Paragraph("Eine randomisierte Pilotstudie mit 48 Studierenden", styles["subtitle"]),
        Spacer(1, 0.5 * cm),
        HRFlowable(width="35%", thickness=2, color=accent, spaceBefore=2, spaceAfter=14, hAlign="LEFT"),
        Paragraph("Fachbereich Psychologie · Musterhochschule · 2026", styles["body"]),
        Paragraph("Abschlussvortrag zur Verteidigung einer Masterarbeit. Alle Inhalte, Zahlen und "
                  "Institutionen sind frei erfunden und dienen nur dem Test des Präsentationstrainers.",
                  styles["small"]),
        PageBreak(),

        Paragraph("1 · Fragestellung", styles["h1"]),
        Paragraph("Untersucht wird, ob eine kurze Pause nach einer Lernphase die Konzentration stärker "
                  "verbessert als eine lange Pause. Die Studie vergleicht eine 5-minütige mit einer "
                  "30-minütigen Pause bei identischer Lernphase.", styles["body"]),
        Paragraph("Hypothese H1", styles["subtitle"]),
        Paragraph("Nach einer 5-minütigen Pause erreichen Teilnehmende im d2-R eine um mindestens 8 Punkte "
                  "höhere Trefferquote als nach einer 30-minütigen Pause.", styles["body"]),
        Paragraph("Hintergrund", styles["subtitle"]),
        Paragraph("Bisherige Untersuchungen berichten uneinheitliche Effekte kurzer Pausen. Viele Studien "
                  "verzichten auf eine randomisierte Zuweisung oder messen die Konzentration erst am "
                  "folgenden Tag. Die vorliegende Pilotstudie prüft daher einen kurzen, standardisierten "
                  "Abstand zwischen Pause und Messung.", styles["body"]),
        Paragraph("Die Studie ist explorativ angelegt. Die Stichprobe erlaubt keine Verallgemeinerung "
                  "auf andere Gruppen oder andere Pausenlängen.", styles["small"]),
        PageBreak(),

        Paragraph("2 · Methode", styles["h1"]),
        Paragraph("Die Zuweisung zu den Gruppen erfolgte per Zufallsliste. Beide Gruppen bearbeiteten "
                  "denselben Text über Arbeitsgedächtnis und hatten dieselbe Lernzeit.", styles["body"]),
        Spacer(1, 0.2 * cm),
        table([
            ["Element", "Umsetzung"],
            ["Stichprobe", "48 Studierende, 24 pro Gruppe"],
            ["Design", "randomisierte Zuweisung, zwei Gruppen"],
            ["Lernphase", "20 Minuten mit identischem Text"],
            ["Pause", "5 bzw. 30 Minuten, medienfrei"],
            ["Messung", "d2-R direkt nach der Pause"],
            ["Auswertung", "t-Test für unabhängige Stichproben, Alpha 0,05"],
        ], [4.2 * cm, 12.4 * cm]),
        Spacer(1, 0.5 * cm),
        Paragraph("Eingeschlossen wurden Studierende ohne diagnostizierte Aufmerksamkeitsstörung. "
                  "Die Teilnahme war freiwillig und unbezahlt. Es gab keine Vorregistrierung.", styles["body"]),
        PageBreak(),

        Paragraph("3 · Ergebnisse", styles["h1"]),
        Paragraph("Die 5-Minuten-Gruppe erreichte im Mittel 62,4 Punkte (SD 10,6), die 30-Minuten-Gruppe "
                  "55,3 Punkte (SD 11,4).", styles["body"]),
        KeepTogether([
            table([
                ["Gruppe", "n", "Mittelwert", "SD"],
                ["5 Minuten Pause", "24", "62,4", "10,6"],
                ["30 Minuten Pause", "24", "55,3", "11,4"],
            ], [6.0 * cm, 2.4 * cm, 4.1 * cm, 4.1 * cm], align_right=True),
            Spacer(1, 0.45 * cm),
            table([
                ["Kennwert", "Wert"],
                ["Differenz der Mittelwerte", "7,1 Punkte"],
                ["t-Test für unabhängige Stichproben", "p = 0,03"],
                ["Effektstärke Cohen's d", "0,65"],
                ["95-%-Konfidenzintervall", "0,7 bis 13,5 Punkte"],
            ], [9.0 * cm, 7.6 * cm]),
        ]),
        Spacer(1, 0.55 * cm),
        Paragraph("Der Unterschied ist statistisch auffällig, verfehlt aber die vorab gesetzte Schwelle "
                  "von 8 Punkten. Die Hypothese H1 wird daher knapp nicht bestätigt.", styles["body"]),
        Paragraph("Ein Teil der Varianz bleibt unerklärt. Die Effektstärke liegt im mittleren Bereich.", styles["small"]),
        PageBreak(),

        Paragraph("4 · Limitationen", styles["h1"]),
        Paragraph("Die Stichprobe umfasst 48 Studierende einer Hochschule. Eine Verallgemeinerung auf "
                  "Berufstätige, Schülerinnen und Schüler oder andere Altersgruppen ist nicht zulässig.", styles["body"]),
        Paragraph("Die Testleitung war nicht geblindet und kannte die Gruppenzuweisung. Es gab nur eine "
                  "Messung direkt nach der Pause, keine Follow-up-Messung.", styles["body"]),
        Paragraph("Die Pause war medienfrei. Ob die Befunde auf Pausen mit Smartphone-Nutzung übertragen "
                  "werden können, ist offen. Die Studie wurde nicht vorregistriert.", styles["body"]),
        Paragraph("Die Ergebnisse beruhen auf einer Gelegenheitsstichprobe aus einem einzigen Semester.", styles["small"]),
        PageBreak(),

        Paragraph("5 · Fazit und Ausblick", styles["h1"]),
        Paragraph("In dieser Stichprobe war die kurze Pause mit 7,1 Punkten mehr verbunden. Der Unterschied "
                  "war statistisch auffällig (p = 0,03), die vorab gesetzte Schwelle von 8 Punkten wurde "
                  "jedoch nicht erreicht.", styles["body"]),
        Paragraph("Die Befunde stützen keine allgemeine Empfehlung für eine bestimmte Pausenlänge. Sie "
                  "sprechen dafür, kurze Pausen in weiteren Studien systematisch zu prüfen.", styles["body"]),
        Spacer(1, 0.3 * cm),
        Paragraph("Nächster Schritt", styles["subtitle"]),
        Paragraph("Geplant ist eine Feldstudie mit 120 Teilnehmenden und zwei Messzeitpunkten. Zusätzlich "
                  "sollen Schlafqualität und Vorwissen erhoben werden.", styles["body"]),
        Paragraph("Diese Planung ist noch nicht umgesetzt und nicht finanziert.", styles["small"]),
    ]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# --------------------------------------------------------------------------- DOCX

def build_docx(path: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    green = RGBColor(0x14, 0x6B, 0x45)
    dark = RGBColor(0x16, 0x24, 0x1D)
    muted = RGBColor(0x5F, 0x72, 0x68)

    def shade(cell, fill: str) -> None:
        properties = cell._tc.get_or_add_tcPr()
        element = OxmlElement("w:shd")
        element.set(qn("w:fill"), fill)
        properties.append(element)

    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2.4)
    section.top_margin = section.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = dark
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    for name, size in (("Heading 1", 15), ("Title", 24)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = green if name == "Heading 1" else dark

    doc.add_paragraph("SYNTHETISCHE DEMO · NONPROFIT-EINLADUNG").runs[0].font.color.rgb = green
    doc.paragraphs[-1].runs[0].font.size = Pt(9)
    doc.paragraphs[-1].runs[0].font.bold = True
    doc.paragraphs[-1].paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph("Tafel Musterstadt", style="Title")
    title.paragraph_format.space_after = Pt(2)
    subtitle = doc.add_paragraph("Einladung zur Unterstützung")
    subtitle.runs[0].font.size = Pt(15)
    subtitle.runs[0].font.color.rgb = muted
    subtitle.paragraph_format.space_after = Pt(2)
    disclaimer = doc.add_paragraph("Fiktive Organisation, alle Zahlen sind synthetische Beispieldaten.")
    disclaimer.runs[0].font.size = Pt(9)
    disclaimer.runs[0].font.italic = True
    disclaimer.runs[0].font.color.rgb = muted
    disclaimer.paragraph_format.space_after = Pt(14)

    def heading(text: str) -> None:
        paragraph = doc.add_heading(text, level=1)
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(6)

    def body(text: str, after: int = 7) -> None:
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(after)

    def bullets(items: list[str]) -> None:
        for item in items:
            paragraph = doc.add_paragraph(item, style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)

    heading("Wer wir sind")
    body("Die Tafel Musterstadt verteilt seit 2014 Lebensmittel an Menschen mit kleinem Budget. "
         "Wir arbeiten mit zwei Standorten in Musterstadt und Umgebung.")
    body("Unser Ziel ist, dass niemand in Musterstadt Lebensmittel wegwerfen muss, während andere zu wenig haben. "
         "Wir holen überschüssige Ware bei Partnermärkten ab und geben sie zweimal pro Woche aus.")

    heading("Was wir 2025 erreicht haben")
    rows = [
        ["Kennzahl", "Wert"],
        ["Haushalte pro Woche", "1.200"],
        ["Ausgegebene Lebensmittel pro Monat", "14 Tonnen"],
        ["Ehrenamtliche", "90"],
        ["Standorte", "2"],
        ["Neu aufgenommene Haushalte 2025", "310"],
        ["Wartezeit auf einen Platz", "6 Wochen"],
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.width = Cm(9.0 if column_index == 0 else 5.5)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.name = "Calibri"
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if row_index == 0:
            for cell in table.rows[0].cells:
                shade(cell, "146B45")
        elif row_index % 2 == 0:
            for cell in table.rows[row_index].cells:
                shade(cell, "EAF5EF")
    doc.add_paragraph()
    body("Im Jahr 2025 haben wir pro Woche 1.200 Haushalte versorgt und dabei 14 Tonnen Lebensmittel pro Monat "
         "ausgegeben. 90 Ehrenamtliche arbeiten an zwei Standorten. 310 Haushalte haben wir neu aufgenommen; "
         "die Wartezeit beträgt sechs Wochen.")

    heading("Wie wir arbeiten")
    bullets([
        "90 Ehrenamtliche übernehmen Abholung, Sortierung und Ausgabe.",
        "Ein Einkauf kostet uns im Schnitt 2,50 Euro.",
        "Mit 25 Euro finanzieren wir zehn Einkäufe.",
        "Frische Ware wird am Ausgabetag verteilt; was übrig bleibt, geht an eine Partnereinrichtung.",
    ])

    heading("Was wir brauchen")
    body("Für 2026 benötigen wir 48.000 Euro, vor allem für Kühlung und Transport. "
         "Die Nachfrage wächst: Die Wartezeit auf einen Platz beträgt aktuell sechs Wochen.")
    body("Wir suchen 15 weitere Ehrenamtliche für Samstagsschichten. "
         "Sachspenden nehmen wir dienstags von 9 bis 12 Uhr an.")

    heading("So können Sie helfen")
    bullets([
        "Spenden Sie einmalig oder monatlich.",
        "Übernehmen Sie eine Samstagsschicht.",
        "Empfehlen Sie uns an Unternehmen in der Region.",
    ])
    body("Jede Unterstützung bleibt in Musterstadt.")

    heading("Häufige Fragen")
    for question, answer in [
        ("Wie viele Menschen erreichen Sie?",
         "Pro Woche versorgen wir 1.200 Haushalte an zwei Standorten."),
        ("Ist die Tafel staatlich finanziert?",
         "Nein. Wir finanzieren uns aus Spenden und Mitgliedsbeiträgen."),
        ("Kann ich auch ohne Anmeldung helfen?",
         "Ja, für die Ausgabe braucht es eine kurze Einweisung vor Ort."),
    ]:
        question_paragraph = doc.add_paragraph(question)
        question_paragraph.runs[0].font.bold = True
        question_paragraph.paragraph_format.space_after = Pt(2)
        answer_paragraph = doc.add_paragraph(answer)
        answer_paragraph.paragraph_format.space_after = Pt(9)

    doc.core_properties.title = "Tafel Musterstadt — Einladung zur Unterstützung"
    doc.core_properties.author = "Synthetische Demo"
    doc.core_properties.subject = "Testpräsentation für den Präsentationstrainer"
    doc.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    targets = [
        (OUT / "01-startup-plantsens.pptx", build_pptx),
        (OUT / "02-thesis-kurzpausen.pdf", build_pdf),
        (OUT / "03-nonprofit-tafel-musterstadt.docx", build_docx),
    ]
    for path, builder in targets:
        builder(path)
        print(f"wrote {path} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
